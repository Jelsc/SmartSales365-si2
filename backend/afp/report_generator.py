"""
Generador de reportes PDF y Excel para el módulo AFP
"""
from io import BytesIO
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter, A4
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, PageBreak
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from decimal import Decimal
from .models import PeriodoFinanciero, BalanceGeneral, EstadoResultados, FlujoEfectivo
from .services import CalculadoraRatios, AnalizadorPatrimonial, EvaluadorAlertas


class GeneradorReportesAFP:
    """Generador de reportes PDF y Excel para AFP"""
    
    @staticmethod
    def generar_pdf(periodo: PeriodoFinanciero) -> BytesIO:
        """Genera un reporte PDF del informe ejecutivo"""
        buffer = BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=A4)
        story = []
        styles = getSampleStyleSheet()
        
        # Estilos personalizados
        title_style = ParagraphStyle(
            'CustomTitle',
            parent=styles['Heading1'],
            fontSize=18,
            textColor=colors.HexColor('#1e40af'),
            spaceAfter=12,
            alignment=TA_CENTER
        )
        
        heading_style = ParagraphStyle(
            'CustomHeading',
            parent=styles['Heading2'],
            fontSize=14,
            textColor=colors.HexColor('#1e40af'),
            spaceAfter=8,
            spaceBefore=12
        )
        
        # Título
        story.append(Paragraph("Informe Ejecutivo Financiero", title_style))
        story.append(Spacer(1, 0.2*inch))
        
        # Información del periodo
        periodo_text = f"Periodo: {periodo.get_tipo_periodo_display()} - {periodo.año}"
        if periodo.mes:
            periodo_text += f" - Mes {periodo.mes}"
        if periodo.trimestre:
            periodo_text += f" - Trimestre {periodo.trimestre}"
        story.append(Paragraph(periodo_text, styles['Normal']))
        story.append(Paragraph(f"Empresa: {periodo.empresa.nombre}", styles['Normal']))
        story.append(Spacer(1, 0.3*inch))
        
        try:
            balance = periodo.balance_general
            estado = periodo.estado_resultados
            flujo = periodo.flujo_efectivo
        except:
            story.append(Paragraph("Error: Faltan estados financieros", styles['Normal']))
            doc.build(story)
            buffer.seek(0)
            return buffer
        
        # Calcular ratios
        balance_anterior = None
        try:
            periodo_anterior = PeriodoFinanciero.objects.filter(
                empresa=periodo.empresa,
                año__lt=periodo.año
            ).order_by('-año', '-mes', '-trimestre').first()
            if periodo_anterior:
                balance_anterior = periodo_anterior.balance_general
        except:
            pass
        
        ratios = CalculadoraRatios.calcular_todos_los_ratios(balance, estado, balance_anterior)
        alertas = EvaluadorAlertas.evaluar_alertas(ratios)
        
        # Resumen de Ratios
        story.append(Paragraph("Resumen de Ratios Clave", heading_style))
        
        ratios_data = [
            ['Ratio', 'Valor'],
            ['Liquidez Corriente', f"{ratios['liquidez']['liquidez_corriente']:.2f}"],
            ['Prueba Ácida', f"{ratios['liquidez']['prueba_acida']:.2f}"],
            ['ROA', f"{ratios['rentabilidad']['roa']*100:.2f}%"],
            ['ROE DuPont', f"{ratios['rentabilidad']['roe_dupont']*100:.2f}%"],
            ['Deuda/Activo', f"{ratios['endeudamiento']['deuda_activo']*100:.2f}%"],
            ['Margen Neto', f"{ratios['rentabilidad']['margen_neto']*100:.2f}%"],
        ]
        
        ratios_table = Table(ratios_data, colWidths=[3*inch, 2*inch])
        ratios_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1e40af')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 12),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
            ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
            ('GRID', (0, 0), (-1, -1), 1, colors.black),
            ('FONTSIZE', (0, 1), (-1, -1), 10),
        ]))
        story.append(ratios_table)
        story.append(Spacer(1, 0.3*inch))
        
        # Alertas
        if alertas:
            story.append(Paragraph("Alertas Financieras", heading_style))
            for alerta in alertas:
                estado_color = colors.red if alerta['estado'] == 'critico' else colors.orange if alerta['estado'] == 'advertencia' else colors.green
                alerta_text = f"<b>{alerta['nombre']}</b>: {alerta['mensaje']} (Valor: {alerta['valor']:.2f})"
                story.append(Paragraph(alerta_text, styles['Normal']))
            story.append(Spacer(1, 0.2*inch))
        
        # Balance General (resumen)
        story.append(PageBreak())
        story.append(Paragraph("Balance General", heading_style))
        
        balance_data = [
            ['Concepto', 'Valor'],
            ['Activo Corriente', f"${balance.activo_corriente:,.2f}"],
            ['Activo No Corriente', f"${balance.activo_no_corriente:,.2f}"],
            ['Total Activo', f"${balance.total_activo:,.2f}"],
            ['Pasivo Corriente', f"${balance.pasivo_corriente:,.2f}"],
            ['Pasivo No Corriente', f"${balance.pasivo_no_corriente:,.2f}"],
            ['Total Pasivo', f"${balance.total_pasivo:,.2f}"],
            ['Patrimonio', f"${balance.total_patrimonio:,.2f}"],
        ]
        
        balance_table = Table(balance_data, colWidths=[3*inch, 2*inch])
        balance_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1e40af')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('ALIGN', (1, 0), (1, -1), 'RIGHT'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 12),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
            ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
            ('GRID', (0, 0), (-1, -1), 1, colors.black),
        ]))
        story.append(balance_table)
        story.append(Spacer(1, 0.3*inch))
        
        # Estado de Resultados (resumen)
        story.append(Paragraph("Estado de Resultados", heading_style))
        
        estado_data = [
            ['Concepto', 'Valor'],
            ['Ventas Netas', f"${estado.ventas_netas:,.2f}"],
            ['Costo de Ventas', f"${estado.costo_ventas:,.2f}"],
            ['Utilidad Bruta', f"${estado.utilidad_bruta:,.2f}"],
            ['Gastos Operativos', f"${estado.gastos_operativos:,.2f}"],
            ['EBIT', f"${estado.ebit:,.2f}"],
            ['Gastos Financieros', f"${estado.gastos_financieros:,.2f}"],
            ['Utilidad Antes de Impuestos', f"${estado.utilidad_antes_impuestos:,.2f}"],
            ['Impuestos', f"${estado.impuestos:,.2f}"],
            ['Utilidad Neta', f"${estado.utilidad_neta:,.2f}"],
        ]
        
        estado_table = Table(estado_data, colWidths=[3*inch, 2*inch])
        estado_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1e40af')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('ALIGN', (1, 0), (1, -1), 'RIGHT'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 12),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
            ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
            ('GRID', (0, 0), (-1, -1), 1, colors.black),
        ]))
        story.append(estado_table)
        
        doc.build(story)
        buffer.seek(0)
        return buffer
    
    @staticmethod
    def generar_excel(periodo: PeriodoFinanciero) -> BytesIO:
        """Genera un reporte Excel del informe ejecutivo"""
        buffer = BytesIO()
        workbook = Workbook()
        
        # Eliminar hoja por defecto
        if 'Sheet' in workbook.sheetnames:
            workbook.remove(workbook['Sheet'])
        
        try:
            balance = periodo.balance_general
            estado = periodo.estado_resultados
            flujo = periodo.flujo_efectivo
        except:
            # Crear hoja de error
            ws = workbook.create_sheet("Error")
            ws['A1'] = "Error: Faltan estados financieros"
            workbook.save(buffer)
            buffer.seek(0)
            return buffer
        
        # Calcular ratios
        balance_anterior = None
        try:
            periodo_anterior = PeriodoFinanciero.objects.filter(
                empresa=periodo.empresa,
                año__lt=periodo.año
            ).order_by('-año', '-mes', '-trimestre').first()
            if periodo_anterior:
                balance_anterior = periodo_anterior.balance_general
        except:
            pass
        
        ratios = CalculadoraRatios.calcular_todos_los_ratios(balance, estado, balance_anterior)
        alertas = EvaluadorAlertas.evaluar_alertas(ratios)
        
        # Estilos
        header_fill = PatternFill(start_color="1e40af", end_color="1e40af", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF", size=12)
        title_font = Font(bold=True, size=14, color="1e40af")
        border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        # Hoja 1: Resumen
        ws_resumen = workbook.create_sheet("Resumen")
        
        ws_resumen['A1'] = "Informe Ejecutivo Financiero"
        ws_resumen['A1'].font = title_font
        ws_resumen.merge_cells('A1:B1')
        
        row = 3
        periodo_text = f"Periodo: {periodo.get_tipo_periodo_display()} - {periodo.año}"
        if periodo.mes:
            periodo_text += f" - Mes {periodo.mes}"
        if periodo.trimestre:
            periodo_text += f" - Trimestre {periodo.trimestre}"
        ws_resumen[f'A{row}'] = periodo_text
        row += 1
        ws_resumen[f'A{row}'] = f"Empresa: {periodo.empresa.nombre}"
        row += 2
        
        # Ratios
        ws_resumen[f'A{row}'] = "Ratios Clave"
        ws_resumen[f'A{row}'].font = title_font
        row += 1
        
        ws_resumen[f'A{row}'] = "Ratio"
        ws_resumen[f'B{row}'] = "Valor"
        for cell in [ws_resumen[f'A{row}'], ws_resumen[f'B{row}']]:
            cell.fill = header_fill
            cell.font = header_font
            cell.border = border
            cell.alignment = Alignment(horizontal='center')
        row += 1
        
        ratios_data = [
            ('Liquidez Corriente', f"{ratios['liquidez']['liquidez_corriente']:.2f}"),
            ('Prueba Ácida', f"{ratios['liquidez']['prueba_acida']:.2f}"),
            ('ROA', f"{ratios['rentabilidad']['roa']*100:.2f}%"),
            ('ROE DuPont', f"{ratios['rentabilidad']['roe_dupont']*100:.2f}%"),
            ('Deuda/Activo', f"{ratios['endeudamiento']['deuda_activo']*100:.2f}%"),
            ('Margen Neto', f"{ratios['rentabilidad']['margen_neto']*100:.2f}%"),
        ]
        
        for ratio_name, ratio_value in ratios_data:
            ws_resumen[f'A{row}'] = ratio_name
            ws_resumen[f'B{row}'] = ratio_value
            for cell in [ws_resumen[f'A{row}'], ws_resumen[f'B{row}']]:
                cell.border = border
            row += 1
        
        # Hoja 2: Balance General
        ws_balance = workbook.create_sheet("Balance General")
        
        ws_balance['A1'] = "Balance General"
        ws_balance['A1'].font = title_font
        ws_balance.merge_cells('A1:B1')
        
        row = 3
        ws_balance[f'A{row}'] = "Concepto"
        ws_balance[f'B{row}'] = "Valor"
        for cell in [ws_balance[f'A{row}'], ws_balance[f'B{row}']]:
            cell.fill = header_fill
            cell.font = header_font
            cell.border = border
            cell.alignment = Alignment(horizontal='center')
        row += 1
        
        balance_data = [
            ('Activo Corriente', balance.activo_corriente),
            ('Activo No Corriente', balance.activo_no_corriente),
            ('Total Activo', balance.total_activo),
            ('Pasivo Corriente', balance.pasivo_corriente),
            ('Pasivo No Corriente', balance.pasivo_no_corriente),
            ('Total Pasivo', balance.total_pasivo),
            ('Patrimonio', balance.total_patrimonio),
        ]
        
        for concepto, valor in balance_data:
            ws_balance[f'A{row}'] = concepto
            ws_balance[f'B{row}'] = f"${float(valor):,.2f}"
            ws_balance[f'B{row}'].alignment = Alignment(horizontal='right')
            for cell in [ws_balance[f'A{row}'], ws_balance[f'B{row}']]:
                cell.border = border
            row += 1
        
        # Hoja 3: Estado de Resultados
        ws_estado = workbook.create_sheet("Estado de Resultados")
        
        ws_estado['A1'] = "Estado de Resultados"
        ws_estado['A1'].font = title_font
        ws_estado.merge_cells('A1:B1')
        
        row = 3
        ws_estado[f'A{row}'] = "Concepto"
        ws_estado[f'B{row}'] = "Valor"
        for cell in [ws_estado[f'A{row}'], ws_estado[f'B{row}']]:
            cell.fill = header_fill
            cell.font = header_font
            cell.border = border
            cell.alignment = Alignment(horizontal='center')
        row += 1
        
        estado_data = [
            ('Ventas Netas', estado.ventas_netas),
            ('Costo de Ventas', estado.costo_ventas),
            ('Utilidad Bruta', estado.utilidad_bruta),
            ('Gastos Operativos', estado.gastos_operativos),
            ('EBIT', estado.ebit),
            ('Gastos Financieros', estado.gastos_financieros),
            ('Utilidad Antes de Impuestos', estado.utilidad_antes_impuestos),
            ('Impuestos', estado.impuestos),
            ('Utilidad Neta', estado.utilidad_neta),
        ]
        
        for concepto, valor in estado_data:
            ws_estado[f'A{row}'] = concepto
            ws_estado[f'B{row}'] = f"${float(valor):,.2f}"
            ws_estado[f'B{row}'].alignment = Alignment(horizontal='right')
            for cell in [ws_estado[f'A{row}'], ws_estado[f'B{row}']]:
                cell.border = border
            row += 1
        
        # Ajustar ancho de columnas
        for ws in [ws_resumen, ws_balance, ws_estado]:
            ws.column_dimensions['A'].width = 30
            ws.column_dimensions['B'].width = 20
        
        workbook.save(buffer)
        buffer.seek(0)
        return buffer

